"""Build a completed test report as an in-memory DOCX file."""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

REPORT_TZ = timezone(timedelta(hours=8))


def clean_text(value):
    return ''.join(c for c in str(value) if c in '\t\n\r' or 0x20 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF)


def upload_name(filename):
    return clean_text((filename or 'Тест.gift').replace('\\', '/').rsplit('/', 1)[-1]).strip()[:200] or 'Тест.gift'


def date_text(value):
    return value.astimezone(REPORT_TZ).strftime('%d.%m.%Y %H:%M:%S') if value else 'Не зафиксировано'


def build_report(data):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[style_name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Title'].font.size = Pt(20)
    doc.styles['Title'].paragraph_format.space_after = Pt(12)
    doc.add_paragraph('Отчёт о проведении тестирования', 'Title')
    p = doc.add_paragraph()
    p.add_run(clean_text(data['test_filename'])).bold = True
    p.paragraph_format.space_after = Pt(12)
    doc.add_paragraph('Тестирование завершено. В отчёте приведены настройки проведения и итоговые результаты участников.')
    metadata = [
        ('Начало', date_text(data['started_at'])),
        ('Окончание', date_text(data['finished_at'])),
        ('Часовой пояс', 'Иркутск UTC+08:00'),
        ('Количество вопросов', data['total_questions']),
        ('Время на вопрос', f"{data['question_seconds']} с"),
        ('Количество участников', len(data['students'])),
    ]
    if data['started_at'] and data['finished_at']:
        seconds = max(0, int((data['finished_at'] - data['started_at']).total_seconds()))
        metadata.insert(2, ('Продолжительность', f'{seconds // 60} мин {seconds % 60} с'))
    for label, value in metadata:
        p = doc.add_paragraph()
        p.add_run(label + ': ').bold = True
        p.add_run(clean_text(value))
    doc.add_heading('Результаты участников', level=1)
    if not data['students']:
        doc.add_paragraph('Участников в проведённом тестировании нет.')
    else:
        table = doc.add_table(rows=1, cols=7)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [0.8, 6.2, 2.0, 1.8, 2.1, 2.5, 2.5]
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        headers = ['№', 'Участник', 'Отвечено', 'Верно', 'Неверно', 'Без ответа', 'Результат']
        for cell, text, width in zip(table.rows[0].cells, headers, widths):
            cell.text = text
            cell.width = Cm(width)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index, row in enumerate(data['students'], start=1):
            cells = table.add_row().cells
            values = [index, row['name'], f"{row['answered']}/{data['total_questions']}", row['correct'], row['incorrect'], row['unanswered'], f"{row['percent']:.1f}%"]
            for cell, value, width in zip(cells, values, widths):
                cell.text = clean_text(value)
                cell.width = Cm(width)
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement('w:' + edge)
            for key, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                el.set(qn('w:' + key), value)
            borders.append(el)
        table._tbl.tblPr.append(borders)
        for index, row in enumerate(table.rows):
            prevent_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(prevent_split)
            for column, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc_pr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'left', 'bottom', 'right'):
                    el = OxmlElement('w:' + edge)
                    el.set(qn('w:w'), '80')
                    el.set(qn('w:type'), 'dxa')
                    margins.append(el)
                tc_pr.append(margins)
                if index == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E9ECEF')
                    tc_pr.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 1 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.bold = index == 0
    note = doc.add_paragraph('Вопрос без отправленного до истечения времени ответа считается неправильным. Столбец «Неверно» включает вопросы из столбца «Без ответа».')
    note.paragraph_format.space_before = Pt(10)
    doc.add_paragraph('Результат в процентах — число правильных ответов, делённое на общее число вопросов. Оценка по учебной шкале не выставляется.')
    doc.core_properties.title = clean_text(data['test_filename'])
    doc.core_properties.subject = 'Отчёт о проведении тестирования'
    doc.core_properties.author = 'Control Test'
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
