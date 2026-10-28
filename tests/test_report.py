from io import BytesIO
import unittest

from docx import Document
from fastapi.testclient import TestClient
from app import main
from app.report import upload_name


class ReportTests(unittest.TestCase):
    def setUp(self):
        main.state = main.AppState()

    def test_report_is_unavailable_until_test_finishes(self):
        with TestClient(main.app) as client:
            response = client.get('/api/results/report.docx')
            self.assertEqual(response.status_code, 409)
            self.assertIn('после завершения', response.json()['error'])

    def test_completed_report_contains_filename_and_disconnected_participant(self):
        with TestClient(main.app) as client:
            response = client.post('/api/upload_test', files={'file': ('Контрольная.gift', '::Первый::Один? {=Да ~Нет}\n\n::Второй::Два? {TRUE}'.encode(), 'text/plain')})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(main.state.test_filename, 'Контрольная.gift')
            main.state.students['old'] = {'name': 'Предыдущая группа', 'ws': None}
            with client.websocket_connect('/ws/student?name=Иванов') as student:
                sid = student.receive_json()['student_id']
                with client.websocket_connect('/ws/teacher') as teacher:
                    teacher.receive_json()
                    teacher.send_json({'type': 'start_test', 'force': True})
                    first = student.receive_json()
                    student.send_json({'type': 'answer', 'question_id': first['question']['id'], 'value': [0]})
                    self.assertEqual(student.receive_json()['type'], 'answer_saved')
                    teacher.send_json({'type': 'next_question', 'current_index': 0})
                    self.assertEqual(student.receive_json()['type'], 'question')
                    teacher.send_json({'type': 'next_question', 'current_index': 1})
                    self.assertEqual(student.receive_json()['type'], 'finished')
            self.assertIsNone(main.state.students[sid]['ws'])
            response = client.get('/api/results/report.docx')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers['content-type'], 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            self.assertIn("filename*=UTF-8''", response.headers['content-disposition'])
            self.assertEqual(response.headers['cache-control'], 'no-store')
            document = Document(BytesIO(response.content))
            self.assertIn('Контрольная.gift', [p.text for p in document.paragraphs])
            rows = document.tables[0].rows
            self.assertEqual(len(rows), 2)
            self.assertEqual([cell.text for cell in rows[1].cells], ['1', 'Иванов', '1/2', '1', '1', '1', '50.0%'])
            self.assertIsNotNone(main.state.started_at)
            self.assertIsNotNone(main.state.finished_at)

    def test_filename_uses_basename_and_strips_xml_controls(self):
        self.assertEqual(upload_name(r'C:\fakepath\Математика.gift'), 'Математика.gift')
        self.assertEqual(upload_name('bad\x01.gift'), 'bad.gift')
